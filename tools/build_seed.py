"""Membuat data awal (seed) prototipe dari file Excel resmi + data dummy.

Pemakaian:
    python3 tools/build_seed.py <Daftar_Indikator_General.xlsx> <template_rb_tematik_2026.xlsx>

Hasil: assets/data/seed.json
- Indikator RB General (27) dan tema/indikator RB Tematik: data asli dari Excel.
- Permasalahan & rencana aksi Tematik: data asli dari template KemenPANRB.
- Sub kegiatan RB General, realisasi, dan tag General: DUMMY (format Sinergi belum diriset).
"""
import json
import random
import sys
from pathlib import Path

import openpyxl

random.seed(2026)
ROOT = Path(__file__).resolve().parent.parent

PD = [
    ("biro-org", "Biro Organisasi", "Biro Organisasi Sekretariat Daerah Provinsi Kalimantan Selatan"),
    ("inspektorat", "Inspektorat", "Inspektorat Provinsi Kalimantan Selatan"),
    ("bappeda", "Bappeda", "Badan Perencanaan Pembangunan Daerah Provinsi Kalimantan Selatan"),
    ("diskominfo", "Diskominfo", "Dinas Komunikasi dan Informatika Provinsi Kalimantan Selatan"),
    ("dispersip", "Dispersip", "Dinas Perpustakaan dan Kearsipan Provinsi Kalimantan Selatan"),
    ("biro-hukum", "Biro Hukum", "Biro Hukum Sekretariat Daerah Provinsi Kalimantan Selatan"),
    ("biro-pbj", "Biro PBJ", "Biro Pengadaan Barang dan Jasa Sekretariat Daerah Provinsi Kalimantan Selatan"),
    ("bkd", "BKD", "Badan Kepegawaian Daerah Provinsi Kalimantan Selatan"),
    ("dpmptsp", "DPMPTSP", "Dinas Penanaman Modal dan Pelayanan Terpadu Satu Pintu"),
    ("disperin", "Dinas Perindustrian", "Dinas Perindustrian"),
    ("dinkes", "Dinas Kesehatan", "Dinas Kesehatan"),
    ("distanpangan", "Distan & KP", "Dinas Pertanian dan Ketahanan Pangan"),
    ("disdikbud", "Disdikbud", "Dinas Pendidikan dan Kebudayaan"),
]
NAME2ID = {full: pid for pid, _s, full in PD}
NAME2ID.update({
    "Biro Organisasi Sekretariat Daerah Provinsi Kalimantan Selatan": "biro-org",
    "Dinas Perpustakaan dan Kearsipan Provinsi Kalimantan Selatan": "dispersip",
})

SUBKOMP = {
    "Strategi Pelaksanaan RB General": "Strategi Pelaksanaan RB General",
    "Capaian Pelaksanaan Kebijakan Reformasi Birokrasi": "Capaian Pelaksanaan Kebijakan Reformasi Birokrasi",
    "Capaian Sasaran Strategis Reformasi Birokrasi": "Capaian Sasaran Strategis Reformasi Birokrasi",
}

TEMA_PENGAMPU = {
    "Pengentasan Kemiskinan": ["bappeda"],
    "Peningkatan Investasi": ["dpmptsp"],
    "Mendorong Hilirisasi": ["bappeda", "disperin"],
    "Mendorong Peningkatan Kualitas dan Akses Layanan Kesehatan": ["dinkes"],
    "Mendukung Ketahanan Pangan Nasional": ["distanpangan"],
    "Mendorong Peningkatan Akses, Kualitas, dan Mutu Pendidikan": ["disdikbud"],
}

# Sub kegiatan DUMMY untuk RB General (nama mengikuti pola nomenklatur, kode fiktif).
SUBKEG_GENERAL = {
    "biro-org": [
        ("Penataan Kelembagaan Perangkat Daerah", "Jumlah dokumen hasil evaluasi kelembagaan", "Dokumen", [0, 1, 0, 1], 310_000_000),
        ("Penyusunan Rencana Aksi dan Evaluasi Reformasi Birokrasi", "Jumlah dokumen rencana aksi dan laporan RB", "Dokumen", [1, 0, 1, 1], 425_000_000),
        ("Fasilitasi Peningkatan Kualitas Pelayanan Publik", "Jumlah unit layanan yang dievaluasi", "Unit", [5, 10, 10, 5], 520_000_000),
        ("Penguatan Akuntabilitas Kinerja Perangkat Daerah", "Jumlah dokumen pengukuran dan pelaporan kinerja", "Dokumen", [1, 1, 1, 1], 380_000_000),
        ("Penyederhanaan Proses Bisnis dan Tata Laksana", "Jumlah proses bisnis yang disederhanakan", "Proses Bisnis", [0, 5, 5, 5], 240_000_000),
        ("Internalisasi Nilai BerAKHLAK", "Jumlah perangkat daerah yang mengikuti internalisasi", "Perangkat Daerah", [10, 10, 10, 11], 180_000_000),
        ("Survei Kepuasan Masyarakat Terpadu", "Jumlah laporan survei kepuasan masyarakat", "Laporan", [0, 1, 0, 1], 290_000_000),
    ],
    "inspektorat": [
        ("Pengawasan Penyelenggaraan Pemerintahan Daerah", "Jumlah laporan hasil pengawasan", "Laporan", [10, 15, 15, 10], 1_250_000_000),
        ("Penilaian Maturitas SPIP", "Jumlah dokumen penilaian maturitas SPIP", "Dokumen", [0, 0, 1, 0], 210_000_000),
        ("Pendampingan Pembangunan Zona Integritas", "Jumlah unit kerja yang didampingi ZI", "Unit Kerja", [0, 5, 5, 0], 275_000_000),
        ("Evaluasi Akuntabilitas Kinerja Internal", "Jumlah laporan hasil evaluasi AKIP", "Laporan", [0, 0, 0, 1], 260_000_000),
        ("Pemantauan Tindak Lanjut Hasil Pemeriksaan", "Jumlah laporan pemantauan tindak lanjut", "Laporan", [1, 1, 1, 1], 190_000_000),
    ],
    "bappeda": [
        ("Penyusunan Dokumen Perencanaan Perangkat Daerah", "Jumlah dokumen perencanaan yang disusun", "Dokumen", [1, 1, 0, 1], 450_000_000),
        ("Koordinasi Penyusunan Perencanaan Kinerja Daerah", "Jumlah dokumen perencanaan kinerja daerah", "Dokumen", [1, 0, 0, 1], 330_000_000),
        ("Monitoring dan Evaluasi Capaian IKU Daerah", "Jumlah laporan evaluasi capaian IKU", "Laporan", [1, 1, 1, 1], 275_000_000),
    ],
    "diskominfo": [
        ("Penatalaksanaan dan Pengawasan e-Government", "Jumlah dokumen arsitektur dan peta rencana SPBE", "Dokumen", [0, 0, 1, 0], 250_000_000),
        ("Pengembangan Aplikasi dan Proses Bisnis Pemerintahan Berbasis Elektronik", "Jumlah aplikasi layanan yang terintegrasi", "Aplikasi", [0, 2, 2, 1], 1_450_000_000),
        ("Pengelolaan Pusat Data Pemerintah Daerah", "Persentase ketersediaan layanan pusat data", "Persen", [99, 99, 99, 99], 820_000_000),
        ("Pengelolaan Pengaduan Masyarakat", "Jumlah laporan tindak lanjut aduan", "Laporan", [1, 1, 1, 1], 180_000_000),
        ("Penyelenggaraan Statistik Sektoral", "Jumlah dokumen data statistik sektoral", "Dokumen", [0, 0, 0, 1], 310_000_000),
        ("Pengelolaan Keamanan Informasi Pemerintahan", "Jumlah layanan keamanan informasi", "Layanan", [1, 1, 1, 1], 410_000_000),
    ],
    "dispersip": [
        ("Pengelolaan Arsip Dinamis dan Alih Media Digital", "Jumlah arsip yang dialihmediakan", "Berkas", [2000, 3000, 3000, 2000], 360_000_000),
        ("Pembinaan Kearsipan Perangkat Daerah", "Jumlah perangkat daerah yang dibina", "Perangkat Daerah", [10, 10, 10, 10], 150_000_000),
    ],
    "biro-hukum": [
        ("Penyusunan dan Harmonisasi Produk Hukum Daerah", "Jumlah produk hukum yang diharmonisasi", "Produk Hukum", [10, 15, 15, 10], 480_000_000),
        ("Analisis dan Evaluasi Kebijakan Daerah", "Jumlah dokumen analisis kebijakan", "Dokumen", [0, 1, 1, 1], 230_000_000),
    ],
    "biro-pbj": [
        ("Pengelolaan Layanan Pengadaan Secara Elektronik", "Jumlah paket pengadaan yang dikelola", "Paket", [40, 80, 60, 20], 520_000_000),
        ("Pembinaan Pelaku Pengadaan", "Jumlah pelaku pengadaan yang dibina", "Orang", [50, 50, 50, 50], 170_000_000),
    ],
    "bkd": [
        ("Penilaian Kompetensi Aparatur", "Jumlah ASN yang dinilai kompetensinya", "Orang", [150, 200, 200, 100], 640_000_000),
        ("Pengelolaan Sistem Informasi Kepegawaian", "Jumlah layanan kepegawaian berbasis elektronik", "Layanan", [2, 2, 2, 2], 390_000_000),
        ("Penerapan Sistem Merit dalam Manajemen ASN", "Jumlah dokumen penerapan sistem merit", "Dokumen", [0, 1, 0, 1], 210_000_000),
    ],
}

# Indikator General → PD : indeks sub kegiatan dummy yang di-tag + status.
TAG_GENERAL = {
    1: [("biro-org", 1, "disetujui")],
    2: [("biro-org", 1, "disetujui")],
    3: [("biro-org", 4, "disetujui"), ("biro-org", 0, "diajukan")],
    4: [("inspektorat", 1, "disetujui"), ("inspektorat", 0, "disetujui")],
    5: [("inspektorat", 2, "disetujui")],
    6: [("bappeda", 1, "disetujui", "Perencanaan Kinerja"), ("biro-org", 3, "disetujui", "Pengukuran Kinerja"),
        ("inspektorat", 3, "diajukan", "Evaluasi Akuntabilitas Kinerja Internal")],
    7: [("bappeda", 0, "disetujui")],
    8: [("diskominfo", 0, "disetujui")],
    9: [("dispersip", 0, "disetujui"), ("dispersip", 1, "diajukan")],
    10: [("diskominfo", 3, "disetujui")],
    11: [("biro-hukum", 1, "disetujui")],
    12: [("biro-hukum", 0, "disetujui")],
    13: [("diskominfo", 4, "diajukan")],
    14: [("biro-pbj", 0, "disetujui"), ("biro-pbj", 1, "disetujui")],
    15: [("bkd", 2, "disetujui"), ("bkd", 0, "diajukan")],
    16: [("biro-org", 2, "disetujui")],
    17: [("biro-org", 2, "disetujui")],
    18: [("diskominfo", 0, "disetujui"), ("diskominfo", 1, "diajukan"), ("diskominfo", 5, "dikembalikan")],
    19: [("inspektorat", 4, "disetujui")],
    20: [("inspektorat", 4, "disetujui")],
    21: [("biro-org", 5, "disetujui")],
    22: [("inspektorat", 0, "diajukan")],
    23: [("biro-org", 6, "disetujui")],
    24: [("bappeda", 2, "disetujui", "[Kelompok IKU makro 1]")],
    27: [("diskominfo", 1, "disetujui")],
}

PEMBAGIAN = {
    6: [
        {"bagian": "Perencanaan Kinerja", "bobot": 30, "pdId": "bappeda"},
        {"bagian": "Pengukuran Kinerja", "bobot": 30, "pdId": "biro-org"},
        {"bagian": "Pelaporan Kinerja", "bobot": 15, "pdId": "biro-org"},
        {"bagian": "Evaluasi Akuntabilitas Kinerja Internal", "bobot": 25, "pdId": "inspektorat"},
    ],
    24: [
        {"bagian": "[Kelompok IKU makro 1]", "bobot": None, "pdId": "bappeda"},
        {"bagian": "[Kelompok IKU makro 2]", "bobot": None, "pdId": "biro-org"},
        {"bagian": "[Kelompok IKU makro 3]", "bobot": None, "pdId": "inspektorat"},
    ],
}


def num(v):
    """Angka dari Excel: dukung teks berkoma ("2,5")."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if "," in str(v) else float(v)
    except ValueError:
        return None


def realisasi_dummy(target):
    """Realisasi output per TW. None = belum dilaporkan. TW IV belum berjalan."""
    out = []
    for i, t in enumerate(target):
        if i == 3 or random.random() < 0.12:
            out.append(None)
        else:
            ratio = random.choice([1, 1, 1, 0.9, 0.8, 1.1, 0.5])
            out.append(round(t * ratio, 2) if t else 0)
    return out


def realisasi_anggaran_dummy(target, pagu, realisasi):
    """Realisasi anggaran per TW (bukan kumulatif). None bila output TW itu belum dilaporkan."""
    tot = sum(target)
    out = []
    for i in range(4):
        if realisasi[i] is None:
            out.append(None)
        else:
            share = (target[i] / tot) if tot else 0.25
            out.append(round(pagu * share * random.uniform(0.55, 1.0), -3))
    return out


def main(path_general, path_tematik):
    data = {"meta": {"tahun": 2026, "dibuat": "seed prototipe", "periodeData": "s.d. Triwulan III"}}
    data["pd"] = [{"id": pid, "singkatan": s, "nama": full} for pid, s, full in PD]

    # ---------- Indikator General (asli) ----------
    ws = openpyxl.load_workbook(path_general, data_only=True).active
    indikator = []
    for r in range(9, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if not isinstance(no, int):
            continue
        koord_raw = ws.cell(r, 8).value or ""
        pengampu = []
        for part in koord_raw.replace(";", "\n").split("\n"):
            part = part.strip()
            if part and NAME2ID.get(part):
                pengampu.append(NAME2ID[part])
        indikator.append({
            "id": f"G{no:02d}",
            "jenis": "general",
            "no": no,
            "komponen": ws.cell(r, 3).value,
            "kelompok": SUBKOMP.get(ws.cell(r, 4).value, ws.cell(r, 4).value),
            "nama": ws.cell(r, 5).value,
            "satuan": "",
            "target": num(ws.cell(r, 6).value),
            "baseline": num(ws.cell(r, 7).value),
            "pengampu": pengampu,
            "pembagian": PEMBAGIAN.get(no, []),
        })

    # ---------- RB Tematik (asli dari template) ----------
    wt = openpyxl.load_workbook(path_tematik, data_only=True)
    sh = wt["DATA RB TEMATIK"]
    header = [c.value for c in sh[1]]
    rows = [dict(zip(header, [c.value for c in row])) for row in sh.iter_rows(min_row=2) if row[0].value]
    tema_list = [c.value for c in wt["MASTER TEMA"]["A"][1:] if c.value]
    data["masterFokus"] = [c.value for c in wt["MASTER FOKUS INTERVENSI"]["A"][1:] if c.value]
    data["masterTema"] = tema_list
    data["templateMeta"] = {r[0].value: r[1].value for r in wt["_META"].iter_rows()}
    data["templatePetunjuk"] = [c.value for c in wt["PETUNJUK"]["A"]]
    data["templateTematikKolom"] = header

    roadmap = {}
    for row in rows:
        key = (row["tema"], row["indikator_roadmap"])
        if key not in roadmap:
            roadmap[key] = row
    for n, (key, row) in enumerate(sorted(roadmap.items(), key=lambda kv: tema_list.index(kv[0][0])), 1):
        tema = key[0]
        indikator.append({
            "id": f"T{n:02d}",
            "jenis": "tematik",
            "no": tema_list.index(tema) + 1,
            "tema": tema,
            "kelompok": tema,
            "sasaran": row["sasaran"],
            "nama": row["indikator_roadmap"],
            "satuan": row["satuan_target"],
            "target": num(row["target"]),
            "baseline": None,
            "pengampu": TEMA_PENGAMPU[tema],
            "pembagian": [],
            "portalId": {"sasaran_roadmap_id": row["sasaran_roadmap_id"], "indikator_roadmap_id": row["indikator_roadmap_id"]},
        })
    data["indikator"] = indikator
    rid = {(i["tema"], i["nama"]): i["id"] for i in indikator if i["jenis"] == "tematik"}

    # ---------- Hak akses ----------
    akses = []
    for ind in indikator:
        for pid in ind["pengampu"]:
            akses.append({
                "id": f"{pid}|{ind['id']}", "pdId": pid, "indikatorId": ind["id"],
                "isiTarget": len(ind["pengampu"]) == 1,
                "tag": True,
                "isiTematik": ind["jenis"] == "tematik",
            })
    data["akses"] = akses

    # ---------- Sub kegiatan + tag ----------
    subkeg, tags, permasalahan = [], [], []
    seq = 1
    for pid, items in SUBKEG_GENERAL.items():
        for j, (nama, out, sat, tgt, pagu) in enumerate(items):
            real = realisasi_dummy(tgt)
            subkeg.append({
                "id": f"{pid}-{j}", "kode": f"9.{seq:02d}.01.2.01.{j + 1:04d}", "pdId": pid, "nama": nama,
                "indikatorOutput": out, "satuan": sat, "target": tgt, "pagu": pagu,
                "realisasi": real, "realisasiAnggaran": realisasi_anggaran_dummy(tgt, pagu, real), "dummy": True,
            })
        seq += 1

    tanggal = ["2026-10-20", "2026-10-22", "2026-10-24", "2026-10-26"]
    tid = 1
    for no, lst in TAG_GENERAL.items():
        for entry in lst:
            pid, idx, status = entry[:3]
            bagian = entry[3] if len(entry) > 3 else ""
            tags.append({
                "id": f"tag{tid:03d}", "jenis": "general", "indikatorId": f"G{no:02d}", "pdId": pid,
                "subkegiatanId": f"{pid}-{idx}", "bagian": bagian, "narasi": "", "anggaranMode": "pagu",
                "anggaranManual": None, "status": status,
                "catatan": "Kaitan dengan indikator perlu dijelaskan di narasi rencana aksi." if status == "dikembalikan" else "",
                "permasalahanId": None, "fokus": None, "diajukan": random.choice(tanggal),
            })
            tid += 1

    # Tematik: hanya pengampu yang disepakati (Hilirisasi sementara Bappeda + Disperin).
    pmap = {}
    for k, row in enumerate(rows):
        pid = NAME2ID.get(row["koordinator"] or "")
        if row["rencana_aksi"] is None or pid not in TEMA_PENGAMPU.get(row["tema"], []):
            continue
        ind_id = rid[(row["tema"], row["indikator_roadmap"])]
        pkey = (ind_id, row["permasalahan_id"])
        if pkey not in pmap:
            pmap[pkey] = f"pm{len(pmap) + 1:03d}"
            permasalahan.append({
                "id": pmap[pkey], "indikatorId": ind_id, "pdId": pid,
                "permasalahan": row["permasalahan"], "sasaran": row["sasaran_permasalahan"],
                "indikator": row["indikator_permasalahan"], "target": num(row["target_indikator_permasalahan"]),
                "satuan": row["satuan_target_indikator_permasalahan"], "realisasi": None, "catatan": "",
                "portalId": {"permasalahan_id": row["permasalahan_id"], "indikator_permasalahan_id": row["indikator_permasalahan_id"]},
            })
        tgt = [num(row[f"target_tw{i}"]) or 0 for i in range(1, 5)]
        real = realisasi_dummy(tgt)
        pagu_t = num(row["anggaran"]) or 0
        sid = f"tm-{k:03d}"
        subkeg.append({
            "id": sid, "kode": f"9.90.01.2.01.{k + 1:04d}", "pdId": pid,
            "nama": "[Dummy] Sub kegiatan pendukung: " + (row["rencana_aksi"] or "")[:70].split("\n")[0],
            "indikatorOutput": row["indikator_output"], "satuan": row["satuan_output"], "target": tgt,
            "pagu": pagu_t, "realisasi": real,
            "realisasiAnggaran": realisasi_anggaran_dummy(tgt, pagu_t, real), "dummy": True,
        })
        tags.append({
            "id": f"tag{tid:03d}", "jenis": "tematik", "indikatorId": ind_id, "pdId": pid, "subkegiatanId": sid,
            "bagian": "", "narasi": row["rencana_aksi"], "anggaranMode": "pagu", "anggaranManual": None,
            "status": "disetujui", "catatan": "", "permasalahanId": pmap[pkey], "fokus": row["fokus_intervensi"],
            "diajukan": "2026-10-15",
            "portalId": {"rencana_aksi_id": row["rencana_aksi_id"], "rencana_aksi_output_id": row["rencana_aksi_output_id"]},
        })
        tid += 1

    data["subkegiatan"] = subkeg
    data["tag"] = tags
    data["permasalahan"] = permasalahan
    data["infografis"] = [
        {"id": "ig1", "judul": "Capaian RB Triwulan III 2026", "kategori": "Rilis capaian", "tanggal": "2026-10-20", "alt": "Contoh infografis capaian RB triwulan III", "gambar": "", "tampil": True},
        {"id": "ig2", "judul": "Apa itu RB General dan RB Tematik?", "kategori": "Penjelasan", "tanggal": "2026-09-02", "alt": "Contoh infografis penjelasan RB", "gambar": "", "tampil": True},
        {"id": "ig3", "judul": "Langkah Pemprov menurunkan kemiskinan", "kategori": "RB Tematik", "tanggal": "2026-08-15", "alt": "Contoh infografis RB Tematik kemiskinan", "gambar": "", "tampil": True},
    ]
    data["imporLog"] = [
        {"id": "log1", "tanggal": "2026-10-14", "periode": "Realisasi s.d. Triwulan III", "baris": len(subkeg), "oleh": "Admin (seed)"},
    ]

    out = ROOT / "assets" / "data" / "seed.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print("indikator", len(indikator), "subkeg", len(subkeg), "tag", len(tags), "permasalahan", len(permasalahan), "->", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
